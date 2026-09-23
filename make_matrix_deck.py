"""
Matrix Construction — Sales Deck Generator
Run: python3 make_matrix_deck.py
Requires: pip install python-pptx
Output: matrix_construction_deck.pptx (same folder)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ── PALETTE ────────────────────────────────────────────────────────────────
VOID        = RGBColor(0x06, 0x06, 0x08)
NAVY        = RGBColor(0x0D, 0x0D, 0x1A)
AMBER       = RGBColor(0xF5, 0xA6, 0x23)
GREEN       = RGBColor(0x00, 0xC8, 0x96)
OFF_WHITE   = RGBColor(0xEB, 0xEB, 0xEB)
MUTED       = RGBColor(0x88, 0x88, 0x88)
CARD        = RGBColor(0x11, 0x11, 0x18)
RED_ALERT   = RGBColor(0xFF, 0x6B, 0x6B)

W = Inches(13.33)
H = Inches(7.5)

out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "matrix_construction_deck.pptx")

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank_layout = prs.slide_layouts[6]

# ── HELPERS ────────────────────────────────────────────────────────────────

def add_rect(slide, x, y, w, h, fill_color=None, line_color=None, line_width=Pt(0)):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, x, y, w, h,
             font_name="Calibri", font_size=Pt(14), bold=False, italic=False,
             color=OFF_WHITE, align=PP_ALIGN.LEFT, wrap=True, margin=Inches(0)):
    txBox = slide.shapes.add_textbox(x, y, w, h)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = margin
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = font_size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txBox

def add_label(slide, text, x, y, w=Inches(5), color=AMBER):
    add_text(slide, text, x, y, w, Inches(0.3),
             font_size=Pt(10), bold=True, color=color)

def slide_bg(slide, color):
    add_rect(slide, 0, 0, W, H, fill_color=color)

def add_tag(slide, text, x, y):
    add_rect(slide, x, y, Inches(2.4), Inches(0.3),
             fill_color=RGBColor(0x3A,0x28,0x00), line_color=AMBER, line_width=Pt(1))
    add_text(slide, text, x+Inches(0.1), y+Inches(0.04), Inches(2.2), Inches(0.22),
             font_size=Pt(9), bold=True, color=AMBER, align=PP_ALIGN.CENTER)

def stat_block(slide, number, label, x, y):
    add_text(slide, number, x, y, Inches(2.8), Inches(0.8),
             font_name="Cambria", font_size=Pt(44), bold=True,
             color=AMBER, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, y+Inches(0.75), Inches(2.8), Inches(0.3),
             font_size=Pt(10), color=MUTED, align=PP_ALIGN.CENTER)

def progress_bar(slide, label, value_text, pct, x, y, bar_color=AMBER):
    add_text(slide, label, x, y, Inches(3.2), Inches(0.22), font_size=Pt(10), color=OFF_WHITE)
    add_text(slide, value_text, x+Inches(3.3), y, Inches(1.5), Inches(0.22),
             font_size=Pt(10), bold=True, color=bar_color, align=PP_ALIGN.RIGHT)
    add_rect(slide, x, y+Inches(0.25), Inches(4.8), Inches(0.07),
             fill_color=RGBColor(0x2A,0x2A,0x3A))
    add_rect(slide, x, y+Inches(0.25), Inches(4.8)*pct, Inches(0.07), fill_color=bar_color)

def feature_bullet(slide, text, x, y, width=Inches(5.0)):
    add_text(slide, "▸  " + text, x, y, width, Inches(0.3),
             font_size=Pt(11), color=OFF_WHITE)

def pain_card(slide, icon, title, desc, x, y, w=Inches(3.8), h=Inches(1.7)):
    add_rect(slide, x, y, w, h, fill_color=CARD,
             line_color=RGBColor(0x2A,0x2A,0x3A), line_width=Pt(1))
    add_text(slide, icon, x+Inches(0.2), y+Inches(0.15), Inches(0.5), Inches(0.4), font_size=Pt(18))
    add_text(slide, title, x+Inches(0.75), y+Inches(0.18), w-Inches(0.95), Inches(0.3),
             font_size=Pt(12), bold=True, color=OFF_WHITE)
    add_text(slide, desc, x+Inches(0.2), y+Inches(0.55), w-Inches(0.35), Inches(1.05),
             font_size=Pt(10), color=MUTED)

def module_card(slide, icon, name, desc, x, y, w=Inches(2.9), h=Inches(1.85)):
    add_rect(slide, x, y, w, h, fill_color=CARD,
             line_color=RGBColor(0x2A,0x2A,0x3A), line_width=Pt(1))
    add_text(slide, icon, x+Inches(0.18), y+Inches(0.15), Inches(0.5), Inches(0.4), font_size=Pt(18))
    add_text(slide, name, x+Inches(0.72), y+Inches(0.18), w-Inches(0.9), Inches(0.28),
             font_size=Pt(12), bold=True, color=OFF_WHITE)
    add_text(slide, desc, x+Inches(0.18), y+Inches(0.58), w-Inches(0.35), Inches(1.0),
             font_size=Pt(10), color=MUTED)

def footer(slide):
    add_rect(slide, 0, H-Inches(0.45), W, Inches(0.45), fill_color=RGBColor(0x03,0x03,0x05))
    add_text(slide, "MATRIX CONSTRUCTION  ·  Contractor Command Platform",
             Inches(0.5), H-Inches(0.38), W-Inches(1), Inches(0.3),
             font_size=Pt(9), color=MUTED, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — HERO
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank_layout)
slide_bg(s, VOID)
add_rect(s, 0, 0, Inches(0.06), H, fill_color=AMBER)
add_tag(s, "CONTRACTOR COMMAND PLATFORM", Inches(0.7), Inches(1.1))
add_text(s, "YOUR SITE.", Inches(0.7), Inches(1.55), Inches(7.5), Inches(1.05),
         font_name="Cambria", font_size=Pt(72), bold=True, color=OFF_WHITE)
add_text(s, "YOUR MONEY.", Inches(0.7), Inches(2.55), Inches(7.5), Inches(1.05),
         font_name="Cambria", font_size=Pt(72), bold=True, color=AMBER)
add_text(s, "YOUR CONTROL.", Inches(0.7), Inches(3.55), Inches(7.5), Inches(1.05),
         font_name="Cambria", font_size=Pt(72), bold=True, color=OFF_WHITE)
add_text(s,
    "Matrix Construction handles every rupee, every worker, every bill — "
    "from the first payment advance to the final company settlement.",
    Inches(0.7), Inches(4.75), Inches(6.5), Inches(0.9),
    font_size=Pt(14), color=MUTED)

add_rect(s, Inches(8.6), Inches(1.1), Inches(4.2), Inches(5.2),
         fill_color=CARD, line_color=RGBColor(0x2A,0x2A,0x3A), line_width=Pt(1))
add_text(s, "AT A GLANCE", Inches(8.8), Inches(1.3), Inches(3.8), Inches(0.3),
         font_size=Pt(10), bold=True, color=AMBER, align=PP_ALIGN.CENTER)
stat_block(s, "500+",  "Workers Tracked",   Inches(8.6), Inches(1.7))
stat_block(s, "50L+",  "Bills Processed",   Inches(8.6), Inches(2.7))
stat_block(s, "Zero",  "Payment Disputes",  Inches(8.6), Inches(3.7))
stat_block(s, "1 App", "For Everything",    Inches(8.6), Inches(4.7))
footer(s)
s.notes_slide.notes_text_frame.text = "Title slide — core contractor value proposition."

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — PROBLEMS
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank_layout)
slide_bg(s, NAVY)
add_rect(s, 0, 0, Inches(0.06), H, fill_color=AMBER)
add_label(s, "THE PAPER TRAIL PROBLEM", Inches(0.7), Inches(0.3))
add_text(s, "What's Costing You Money Right Now",
         Inches(0.7), Inches(0.6), Inches(12), Inches(0.75),
         font_name="Cambria", font_size=Pt(38), bold=True, color=OFF_WHITE)
add_text(s, "Every contractor faces these daily — most just accept them as the cost of doing business.",
         Inches(0.7), Inches(1.38), Inches(12), Inches(0.35), font_size=Pt(13), color=MUTED)

cards = [
    ("📋", "Register gets lost or tampered",
     "Manual registers go missing or get altered. You pay for days never worked."),
    ("💸", "Advance payments with no trail",
     "Give ₹5,000 on Monday. By month end it's gone — no record anywhere."),
    ("🧾", "Company bills don't match",
     "Company says you've been paid. You disagree. Neither side has clean proof."),
    ("📂", "Proofs scattered on WhatsApp",
     "Screenshots buried in 12 chats. Proving anything takes an hour of scrolling."),
    ("🔢", "Wage calculation errors",
     "Counting half-days on a notebook. One error and you overpay or underpay."),
    ("📊", "No picture of site expenses",
     "You don't know if the site is profitable until the final bill arrives months later."),
]
positions = [
    (Inches(0.7), Inches(1.9)), (Inches(4.65), Inches(1.9)), (Inches(8.6), Inches(1.9)),
    (Inches(0.7), Inches(3.75)), (Inches(4.65), Inches(3.75)), (Inches(8.6), Inches(3.75)),
]
for (icon, title, desc), (x, y) in zip(cards, positions):
    pain_card(s, icon, title, desc, x, y)
footer(s)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — ATTENDANCE
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank_layout)
slide_bg(s, VOID)
add_rect(s, 0, 0, Inches(0.06), H, fill_color=GREEN)
add_label(s, "MODULE 01", Inches(0.7), Inches(0.3), color=GREEN)
add_text(s, "Attendance That Sticks",
         Inches(0.7), Inches(0.6), Inches(6), Inches(0.7),
         font_name="Cambria", font_size=Pt(36), bold=True, color=OFF_WHITE)
add_text(s, "Mark daily attendance in seconds. Full day, half day, absent — timestamped. No paper, no disputes.",
         Inches(0.7), Inches(1.35), Inches(5.8), Inches(0.6), font_size=Pt(13), color=MUTED)
for i, b in enumerate([
    "Worker-wise daily log with date and shift",
    "Bulk mark for large gangs (30+ workers)",
    "Monthly summary auto-generated",
    "Edit with audit trail — nothing disappears quietly",
]):
    feature_bullet(s, b, Inches(0.7), Inches(2.1) + Inches(0.38)*i)

tx, ty, tw = Inches(7.0), Inches(0.45), Inches(5.8)
add_rect(s, tx, ty, tw, Inches(6.65), fill_color=CARD,
         line_color=RGBColor(0x2A,0x2A,0x3A), line_width=Pt(1))
add_rect(s, tx, ty, tw, Inches(0.38), fill_color=RGBColor(0x1A,0x1A,0x28))
cols_x = [tx+Inches(0.15), tx+Inches(1.8), tx+Inches(3.1), tx+Inches(4.3)]
for hx, ht in zip(cols_x, ["WORKER","TRADE","TODAY","DAYS"]):
    add_text(s, ht, hx, ty+Inches(0.07), Inches(1.2), Inches(0.24),
             font_size=Pt(9), bold=True, color=MUTED)
rows = [
    ("Ramu K.","Mason","Full","22.5",GREEN), ("Suresh P.","Helper","Half","19.0",AMBER),
    ("Ajay M.","Welder","Absent","20.0",RED_ALERT), ("Vinod R.","Mason","Full","23.0",GREEN),
    ("Mohan D.","Painter","Full","21.5",GREEN), ("Karan S.","Helper","Full","22.0",GREEN),
    ("Deepak T.","Welder","Half","18.5",AMBER), ("Priya N.","Mason","Full","23.0",GREEN),
    ("Anwar H.","Painter","Absent","17.0",RED_ALERT), ("Sunil B.","Helper","Full","22.5",GREEN),
]
for i, (name, trade, status, days, scol) in enumerate(rows):
    ry = ty + Inches(0.38) + Inches(0.55)*i
    if i % 2 == 0:
        add_rect(s, tx, ry, tw, Inches(0.54), fill_color=RGBColor(0x14,0x14,0x1E))
    add_text(s, name,   cols_x[0], ry+Inches(0.12), Inches(1.5), Inches(0.3), font_size=Pt(11), color=OFF_WHITE)
    add_text(s, trade,  cols_x[1], ry+Inches(0.12), Inches(1.2), Inches(0.3), font_size=Pt(11), color=MUTED)
    add_text(s, status, cols_x[2], ry+Inches(0.12), Inches(1.0), Inches(0.3), font_size=Pt(11), bold=True, color=scol)
    add_text(s, days,   cols_x[3], ry+Inches(0.12), Inches(0.8), Inches(0.3), font_size=Pt(11), color=OFF_WHITE)
footer(s)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — WAGES
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank_layout)
slide_bg(s, NAVY)
add_rect(s, 0, 0, Inches(0.06), H, fill_color=AMBER)
add_label(s, "MODULE 02", Inches(0.7), Inches(0.3))
add_text(s, "Wages & Advances",
         Inches(0.7), Inches(0.6), Inches(6), Inches(0.7),
         font_name="Cambria", font_size=Pt(36), bold=True, color=OFF_WHITE)
add_text(s, "Per-worker rates, auto-deducted advances, half-day calculations — wage sheet built automatically.",
         Inches(0.7), Inches(1.35), Inches(5.8), Inches(0.6), font_size=Pt(13), color=MUTED)
for i, b in enumerate([
    "Per-worker daily rate setup",
    "Pre-payment (advance) auto-deducted from wages",
    "Post-payment tracking after work closes",
    "Printable wage slip per worker",
]):
    feature_bullet(s, b, Inches(0.7), Inches(2.1) + Inches(0.38)*i)

tx, ty = Inches(7.0), Inches(0.75)
add_rect(s, tx, ty, Inches(5.8), Inches(5.9), fill_color=CARD,
         line_color=RGBColor(0x2A,0x2A,0x3A), line_width=Pt(1))
add_rect(s, tx, ty, Inches(5.8), Inches(0.42), fill_color=RGBColor(0x1A,0x1A,0x28))
add_text(s, "Wage Sheet — Ramu K.  |  October 2026",
         tx+Inches(0.18), ty+Inches(0.08), Inches(5.4), Inches(0.26),
         font_size=Pt(11), bold=True, color=MUTED)
wage_rows = [
    ("Days Worked",        "22.5 days",  OFF_WHITE),
    ("Daily Rate",         "₹750 / day", OFF_WHITE),
    ("Gross Wages",        "+ ₹16,875",  GREEN),
    ("Advance Paid (Pre)", "− ₹3,000",   RED_ALERT),
    ("Material Deduction", "− ₹500",     RED_ALERT),
    ("Other Advance",      "− ₹1,000",   RED_ALERT),
]
for i, (lbl, val, vcol) in enumerate(wage_rows):
    wy = ty + Inches(0.52) + Inches(0.7)*i
    if i % 2 == 0:
        add_rect(s, tx, wy, Inches(5.8), Inches(0.68), fill_color=RGBColor(0x14,0x14,0x1E))
    add_text(s, lbl, tx+Inches(0.2), wy+Inches(0.18), Inches(3.2), Inches(0.3),
             font_size=Pt(12), color=MUTED)
    add_text(s, val, tx+Inches(3.5), wy+Inches(0.18), Inches(2.0), Inches(0.3),
             font_size=Pt(12), bold=True, color=vcol, align=PP_ALIGN.RIGHT)
wy = ty + Inches(0.52) + Inches(0.7)*6
add_rect(s, tx, wy, Inches(5.8), Inches(0.75), fill_color=RGBColor(0x1E,0x14,0x00))
add_rect(s, tx, wy, Inches(5.8), Pt(2), fill_color=AMBER)
add_text(s, "NET PAYABLE", tx+Inches(0.2), wy+Inches(0.2), Inches(3), Inches(0.45),
         font_name="Cambria", font_size=Pt(18), bold=True, color=OFF_WHITE)
add_text(s, "₹12,375", tx+Inches(3.4), wy+Inches(0.2), Inches(2.1), Inches(0.45),
         font_name="Cambria", font_size=Pt(18), bold=True, color=AMBER, align=PP_ALIGN.RIGHT)
footer(s)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — COMPANY BILLS
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank_layout)
slide_bg(s, VOID)
add_rect(s, 0, 0, Inches(0.06), H, fill_color=RED_ALERT)
add_label(s, "MODULE 03", Inches(0.7), Inches(0.3), color=RED_ALERT)
add_text(s, "Company Bills & Settlement",
         Inches(0.7), Inches(0.6), Inches(6), Inches(0.7),
         font_name="Cambria", font_size=Pt(36), bold=True, color=OFF_WHITE)
add_text(s, "Track every bill, payment received, and outstanding balance — with proof attachments.",
         Inches(0.7), Inches(1.35), Inches(5.8), Inches(0.6), font_size=Pt(13), color=MUTED)
for i, b in enumerate([
    "Bill raised vs bill approved tracking",
    "Payment received log with dates and references",
    "Attach photos of signed bills and receipts",
    "Running balance: what's owed to you right now",
]):
    feature_bullet(s, b, Inches(0.7), Inches(2.1) + Inches(0.38)*i)

tx, ty = Inches(7.0), Inches(0.65)
add_rect(s, tx, ty, Inches(5.8), Inches(6.1), fill_color=CARD,
         line_color=RGBColor(0x2A,0x2A,0x3A), line_width=Pt(1))
add_rect(s, tx, ty, Inches(5.8), Inches(0.42), fill_color=RGBColor(0x1A,0x1A,0x28))
add_text(s, "Project: Tower B  |  Bill #04",
         tx+Inches(0.18), ty+Inches(0.08), Inches(5.3), Inches(0.26),
         font_size=Pt(11), bold=True, color=MUTED)
py = ty + Inches(0.6)
progress_bar(s, "Bill #04 Raised",  "₹2,40,000", 1.00, tx+Inches(0.2), py,             bar_color=AMBER)
progress_bar(s, "Bill Approved",    "₹2,20,000", 0.92, tx+Inches(0.2), py+Inches(0.72), bar_color=RGBColor(0x4D,0x9F,0xFF))
progress_bar(s, "Payment Received", "₹1,60,000", 0.67, tx+Inches(0.2), py+Inches(1.44), bar_color=GREEN)

add_rect(s, tx, ty+Inches(2.3), Inches(5.8), Pt(1), fill_color=RGBColor(0x2A,0x2A,0x3A))
add_text(s, "OUTSTANDING BALANCE", tx+Inches(0.2), ty+Inches(2.45), Inches(3.5), Inches(0.35),
         font_name="Cambria", font_size=Pt(16), bold=True, color=OFF_WHITE)
add_text(s, "₹60,000", tx+Inches(3.8), ty+Inches(2.45), Inches(1.75), Inches(0.35),
         font_name="Cambria", font_size=Pt(20), bold=True, color=RED_ALERT, align=PP_ALIGN.RIGHT)

add_text(s, "Recent Payments", tx+Inches(0.2), ty+Inches(3.0), Inches(4), Inches(0.28),
         font_size=Pt(10), bold=True, color=AMBER)
for i, (dt, ref, amt) in enumerate([
    ("15 Oct 2026", "RTGS Transfer", "+ ₹80,000"),
    ("02 Oct 2026", "Cheque #4421",  "+ ₹50,000"),
    ("20 Sep 2026", "NEFT Transfer", "+ ₹30,000"),
]):
    ry = ty + Inches(3.35) + Inches(0.73)*i
    add_rect(s, tx+Inches(0.2), ry, Inches(5.3), Inches(0.63), fill_color=RGBColor(0x14,0x14,0x1E))
    add_text(s, dt,  tx+Inches(0.35), ry+Inches(0.08), Inches(1.4), Inches(0.22), font_size=Pt(10), color=MUTED)
    add_text(s, ref, tx+Inches(1.8),  ry+Inches(0.08), Inches(2.0), Inches(0.22), font_size=Pt(10), color=OFF_WHITE)
    add_text(s, amt, tx+Inches(3.85), ry+Inches(0.08), Inches(1.4), Inches(0.22),
             font_size=Pt(11), bold=True, color=GREEN, align=PP_ALIGN.RIGHT)
    add_text(s, "📎 Proof attached", tx+Inches(0.35), ry+Inches(0.35), Inches(2.5), Inches(0.2),
             font_size=Pt(9), color=GREEN)
footer(s)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — HOW IT WORKS
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank_layout)
slide_bg(s, NAVY)
add_rect(s, 0, 0, Inches(0.06), H, fill_color=AMBER)
add_label(s, "SIMPLE BY DESIGN", Inches(0.7), Inches(0.3))
add_text(s, "How a Day Looks with Matrix",
         Inches(0.7), Inches(0.6), Inches(12), Inches(0.75),
         font_name="Cambria", font_size=Pt(38), bold=True, color=OFF_WHITE)
add_text(s, "From site morning to end-of-month settlement — every step in one place.",
         Inches(0.7), Inches(1.4), Inches(12), Inches(0.35), font_size=Pt(13), color=MUTED)

add_rect(s, Inches(1.22), Inches(2.46), Inches(10.8), Pt(2), fill_color=AMBER)
steps = [
    (1, "Mark\nAttendance",   "Tap each worker's\nstatus at site start.\n2 min for 30 workers."),
    (2, "Log\nAdvances",      "Worker asks for early\npay? Log it instantly.\nAuto-deducts from wages."),
    (3, "Record\nExpenses",   "Material, tools,\ntransport — every cost\nwith a photo proof."),
    (4, "Raise\nCompany Bill","Matrix shows your\nexact total. One tap\nto generate the bill."),
    (5, "Log Payment\nReceived","Mark it received,\nattach screenshot.\nBalance updates live."),
]
for i, (num, title, desc) in enumerate(steps):
    cx = Inches(0.85) + Inches(2.5)*i
    circ = s.shapes.add_shape(9, cx, Inches(2.1), Inches(0.72), Inches(0.72))
    circ.fill.solid(); circ.fill.fore_color.rgb = AMBER
    circ.line.fill.background()
    add_text(s, str(num), cx+Inches(0.13), Inches(2.17), Inches(0.46), Inches(0.46),
             font_name="Cambria", font_size=Pt(20), bold=True, color=VOID, align=PP_ALIGN.CENTER)
    add_text(s, title, cx-Inches(0.2), Inches(2.95), Inches(2.2), Inches(0.6),
             font_name="Cambria", font_size=Pt(14), bold=True, color=OFF_WHITE, align=PP_ALIGN.CENTER)
    add_text(s, desc, cx-Inches(0.2), Inches(3.62), Inches(2.2), Inches(1.1),
             font_size=Pt(10), color=MUTED, align=PP_ALIGN.CENTER)

add_rect(s, Inches(0.7), Inches(5.55), Inches(11.93), Inches(1.3),
         fill_color=RGBColor(0x1A,0x12,0x00), line_color=AMBER, line_width=Pt(1))
add_text(s, "💡  End of month: Matrix generates your complete wage sheet, expense summary, and pending bill — all in one report.",
         Inches(0.95), Inches(5.72), Inches(11.4), Inches(0.8), font_size=Pt(13), color=OFF_WHITE)
footer(s)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — ALL MODULES
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank_layout)
slide_bg(s, VOID)
add_rect(s, 0, 0, Inches(0.06), H, fill_color=GREEN)
add_label(s, "FULL PLATFORM", Inches(0.7), Inches(0.3), color=GREEN)
add_text(s, "Every Module You Need",
         Inches(0.7), Inches(0.6), Inches(12), Inches(0.65),
         font_name="Cambria", font_size=Pt(38), bold=True, color=OFF_WHITE)
add_text(s, "Nothing missing. Nothing you don't need.",
         Inches(0.7), Inches(1.28), Inches(7), Inches(0.35), font_size=Pt(13), color=MUTED)

modules = [
    ("📅","Labour Attendance",  "Daily full/half/absent marking.\nMonthly totals auto-calculated."),
    ("💰","Wage Management",    "Per-worker rates, deductions,\nand printable wage slips."),
    ("⬆️", "Pre Payments",       "Advance payments logged &\nauto-adjusted against wages."),
    ("⬇️", "Post Payments",      "Final settlements tracked\nwith full payment history."),
    ("🏢","Company Bills",      "Raise bills, track approvals,\nlog payments received."),
    ("📎","Payment Proofs",     "Attach receipt photos &\nscreenshots to every transaction."),
    ("📊","Site Expenses",      "Track material, tool & transport\ncosts vs project budget."),
    ("📈","Project Ledger",     "Full income vs expense view.\nKnow your site margin anytime."),
]
cw, ch, gap = Inches(2.9), Inches(1.85), Inches(0.32)
sx, sy = Inches(0.7), Inches(1.75)
for i, (icon, name, desc) in enumerate(modules):
    col, row = i % 4, i // 4
    module_card(s, icon, name, desc, sx+(cw+gap)*col, sy+(ch+gap)*row, w=cw, h=ch)
footer(s)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — CTA
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank_layout)
slide_bg(s, VOID)
glow = s.shapes.add_shape(9, Inches(3.5), Inches(0.8), Inches(6.5), Inches(5.8))
glow.fill.solid(); glow.fill.fore_color.rgb = RGBColor(0x18,0x0C,0x00)
glow.line.fill.background()
add_rect(s, 0, 0, Inches(0.06), H, fill_color=AMBER)
add_label(s, "READY TO START?", Inches(0.7), Inches(0.55))
add_text(s, "Stop Losing Money",
         Inches(0.7), Inches(0.95), Inches(8), Inches(0.88),
         font_name="Cambria", font_size=Pt(58), bold=True, color=OFF_WHITE)
add_text(s, "to Paperwork.",
         Inches(0.7), Inches(1.8), Inches(8), Inches(0.88),
         font_name="Cambria", font_size=Pt(58), bold=True, color=AMBER)
add_text(s,
    "Matrix Construction is built for contractors who run real sites "
    "and need answers — not spreadsheets.",
    Inches(0.7), Inches(2.85), Inches(6.5), Inches(0.7),
    font_size=Pt(14), color=MUTED)

add_rect(s, Inches(0.7), Inches(3.7), Inches(2.4), Inches(0.55), fill_color=AMBER)
add_text(s, "Book a Demo", Inches(0.7), Inches(3.77), Inches(2.4), Inches(0.42),
         font_size=Pt(14), bold=True, color=VOID, align=PP_ALIGN.CENTER)
add_rect(s, Inches(3.3), Inches(3.7), Inches(2.4), Inches(0.55),
         fill_color=CARD, line_color=OFF_WHITE, line_width=Pt(1))
add_text(s, "Call Us Now", Inches(3.3), Inches(3.77), Inches(2.4), Inches(0.42),
         font_size=Pt(14), bold=True, color=OFF_WHITE, align=PP_ALIGN.CENTER)

add_rect(s, Inches(0.7), Inches(4.5), Inches(11.93), Pt(1), fill_color=RGBColor(0x2A,0x2A,0x3A))
for i, (key, val) in enumerate([
    ("EMAIL",    "hello@matrixconstruction.in"),
    ("WHATSAPP", "+91 99999 99999"),
    ("WEBSITE",  "matrixconstruction.in"),
]):
    cx = Inches(0.7) + Inches(4.0)*i
    add_text(s, val, cx, Inches(4.68), Inches(3.8), Inches(0.35),
             font_size=Pt(14), bold=True, color=AMBER)
    add_text(s, key, cx, Inches(5.05), Inches(3.8), Inches(0.25),
             font_size=Pt(10), color=MUTED)

add_text(s, "MATRIX\nCONSTRUCTION",
         Inches(9.0), Inches(4.8), Inches(4.2), Inches(2.2),
         font_name="Cambria", font_size=Pt(38), bold=True,
         color=RGBColor(0x1A,0x1A,0x2A), align=PP_ALIGN.RIGHT)

add_rect(s, 0, H-Inches(0.45), W, Inches(0.45), fill_color=RGBColor(0x03,0x03,0x05))
add_text(s, "MATRIX CONSTRUCTION  ·  Contractor Command Platform  ·  © 2026",
         Inches(0.5), H-Inches(0.38), W-Inches(1), Inches(0.3),
         font_size=Pt(9), color=MUTED, align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = "Closing slide — book a demo or WhatsApp contact."

# ── SAVE ────────────────────────────────────────────────────────────────────
prs.save(out_path)
print(f"Saved: {out_path}")
